from io import BytesIO
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from app.ingestion.extraction import (
    DocumentExtractionError,
    UnsupportedDocumentError,
    _markdown_to_evidence_text,
    extract_document,
)


def test_markdown_table_rows_preserve_all_column_relationships():
    content = _markdown_to_evidence_text(
        """## Équipe

| Nom | Rôle | Entité | Charge | Période |
|---|---|---|---|---|
| Fatou Sy | Conduite du changement et formation | Avaliance | 50 % | Sept. 2024 — Juin 2025 |

## Budget

| Poste | Responsable / équipe | Jours | TJM moyen | Montant HT |
|---|---|---|---|---|
| Conduite du changement et formation | Fatou Sy | 112 | 700 € | 78 400 € |
"""
    )

    assert "Nom: Fatou Sy; Rôle: Conduite du changement et formation; Entité: Avaliance; Charge: 50 %; Période: Sept. 2024 — Juin 2025" in content
    assert "Poste: Conduite du changement et formation; Responsable / équipe: Fatou Sy; Jours: 112; TJM moyen: 700 €; Montant HT: 78 400 €" in content


def test_extract_txt_normalizes_utf8_content():
    sections = extract_document("mission.TXT", b"Titre\r\n\r\n  Besoin   cloud  ")

    assert len(sections) == 1
    assert sections[0].content == "Titre\n\nBesoin cloud"
    assert sections[0].page is None


def test_extract_docx_preserves_paragraphs_and_tables():
    document = Document()
    document.add_heading("Mission ERP", level=1)
    document.add_paragraph("Migration et conduite du changement")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Technologie"
    table.cell(0, 1).text = "SAP"
    stream = BytesIO()
    document.save(stream)

    sections = extract_document("mission.docx", stream.getvalue())

    assert len(sections) == 1
    assert sections[0].content == (
        "Mission ERP\n\nMigration et conduite du changement\n\nTechnologie | SAP"
    )


def test_extract_pdf_rejects_page_without_text_as_scanned():
    document = MagicMock()
    document.pages = {1: MagicMock()}
    document.export_to_markdown.return_value = ""
    converter = MagicMock()
    converter.convert.return_value.document = document

    with patch("app.ingestion.extraction.DocumentConverter", return_value=converter), pytest.raises(
        DocumentExtractionError, match="PDF scannés"
    ):
        extract_document("scan.pdf", b"pdf")


def test_extract_pdf_preserves_physical_pages_after_blank_cover():
    document = MagicMock()
    document.pages = {1: MagicMock(), 2: MagicMock(), 3: MagicMock()}
    document.export_to_markdown.side_effect = ["", "", "Contexte de la mission"]
    converter = MagicMock()
    converter.convert.return_value.document = document

    with patch("app.ingestion.extraction.DocumentConverter", return_value=converter):
        sections = extract_document("cover.pdf", b"pdf")

    assert [(section.page, section.page_count) for section in sections] == [(3, 3)]
    assert sections[0].content == "Contexte de la mission"
    assert document.export_to_markdown.call_args_list[-1].kwargs == {
        "page_no": 3,
        "compact_tables": False,
    }


def test_extract_pdf_rejects_document_when_any_physical_page_export_fails():
    document = MagicMock()
    document.pages = {1: MagicMock(), 2: MagicMock(), 3: MagicMock()}
    document.export_to_markdown.side_effect = ["Première page exploitable", ValueError("page malformed")]
    converter = MagicMock()
    converter.convert.return_value.document = document

    with (
        patch("app.ingestion.extraction.DocumentConverter", return_value=converter),
        pytest.raises(DocumentExtractionError, match="page PDF 2"),
    ):
        extract_document("partial.pdf", b"pdf")


def test_extract_pdf_allows_empty_password_decryption():
    document = MagicMock()
    document.pages = {1: MagicMock()}
    document.export_to_markdown.return_value = "Texte PDF valide"
    converter = MagicMock()
    converter.convert.return_value.document = document
    reader = MagicMock(is_encrypted=True)
    reader.decrypt.return_value = 1

    with (
        patch("app.ingestion.extraction.PdfReader", return_value=reader),
        patch("app.ingestion.extraction.DocumentConverter", return_value=converter),
    ):
        sections = extract_document("valid.pdf", b"pdf")

    reader.decrypt.assert_called_once_with("")
    assert sections[0].content == "Texte PDF valide"


def test_extract_pdf_rejects_password_protected_pdf():
    reader = MagicMock(is_encrypted=True)
    reader.decrypt.return_value = 0

    with patch("app.ingestion.extraction.PdfReader", return_value=reader), pytest.raises(
        DocumentExtractionError, match="protégé par mot de passe"
    ):
        extract_document("protected.pdf", b"pdf")


def test_extract_pdf_falls_back_to_pypdf_when_docling_fails():
    reader = MagicMock(is_encrypted=False)
    page = MagicMock()
    page.extract_text.return_value = "Texte sélectionnable issu du PDF"
    reader.pages = [page]

    with (
        patch("app.ingestion.extraction.DocumentConverter", side_effect=ImportError("libxcb.so.1")),
        patch("app.ingestion.extraction.PdfReader", return_value=reader),
    ):
        sections = extract_document("standard.pdf", b"pdf")

    assert [(section.page, section.page_count, section.content) for section in sections] == [
        (1, 1, "Texte sélectionnable issu du PDF")
    ]


def test_extract_rejects_unsupported_extension():
    with pytest.raises(UnsupportedDocumentError, match="PDF, DOCX ou TXT"):
        extract_document("archive.doc", b"legacy")